from django.shortcuts import render, redirect
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.contrib import messages
from django.core.mail import send_mail
from django.conf import settings
from .models import OTPVerification
import random
from django.http import JsonResponse
from .models import Tickets
from .forms import Booking
from .models import Customer
from .forms import CustomerForm
from .models import Agent
from .forms import AgentForm
from django.views.decorators.csrf import csrf_exempt
from django.http import HttpResponse
from openpyxl import Workbook
from .models import Tickets
from io import BytesIO
from openpyxl.styles import Alignment, Font, Protection

def signup_view(request):
    if request.method == "POST":
        email = request.POST.get("email").strip().lower()

        otp = str(random.randint(100000, 999999))

        OTPVerification.objects.update_or_create(
            email=email,
            defaults={"otp": otp}
        )

        # OTP is ALWAYS sent to your fixed email
        otp_receiver = "sashwathisri@gmail.com"

        send_mail(
            "OTP Verification",
            f"Your OTP is {otp}",
            settings.EMAIL_HOST_USER,
            [otp_receiver],
            fail_silently=False,
        )

        # Keep the user's entered email for account creation
        request.session["signup_email"] = email

        return redirect("verify_otp")

    return render(request, "two/signup.html")


def verify_otp(request):
    email = request.session.get("signup_email")
    if not email:
        return redirect("signup")
    if request.method == "POST":
        otp = request.POST.get("otp")
        try:
            record = OTPVerification.objects.get(email=email)
            if record.otp == otp:
                return redirect("set_password")
            else:
                messages.error(request, "Invalid OTP")
        except OTPVerification.DoesNotExist:
            messages.error(request, "OTP verification failed")
    return render(request, "two/verify_otp.html")


def set_password(request):
    email = request.session.get("signup_email")

    if not email:
        return redirect("signup")

    if request.method == "POST":
        password = request.POST.get("password")
        confirm_password = request.POST.get("confirm_password")

        if password != confirm_password:
            messages.error(request, "Passwords do not match")
            return redirect("set_password")

        # Check whether this email already has an account
        user = User.objects.filter(email__iexact=email).first()

        if user:
            # Existing account:
            # Change the password instead of creating a new account
            user.set_password(password)
            user.save()

            messages.success(
                request,
                "Password updated successfully. You can now login."
            )

        else:
            # New account:
            username = email.split("@")[0]

            # Make username unique if necessary
            if User.objects.filter(username=username).exists():
                username = username + str(random.randint(1000, 9999))

            User.objects.create_user(
                username=username,
                email=email,
                password=password
            )

            messages.success(
                request,
                "Account created successfully. You can now login."
            )

        # Remove OTP after successful password setup
        OTPVerification.objects.filter(email=email).delete()

        # Remove email from session
        request.session.pop("signup_email", None)

        return redirect("login")

    return render(request, "two/set_password.html")


def login_view(request):
    if request.method == "POST":

        email = request.POST.get("username", "").strip()
        password = request.POST.get("password", "")

        # Find user using email
        user_obj = User.objects.filter(email__iexact=email).first()

        if user_obj is None:
            messages.error(request, "Email not registered")
            return render(request, "two/login.html")

        # Check password
        user = authenticate(
            request,
            username=user_obj.username,
            password=password
        )

        if user is not None:
            login(request, user)
            return redirect("home")

        messages.error(request, "Invalid Password")

    return render(request, "two/login.html")


def home(request):

    return render(request, "two/home.html")


def logout_view(request):

    logout(request)

    return redirect("login")

def index(request):
    return render(request, 'two/index.html')    

@csrf_exempt
def upload(request):
    if request.method == "POST":
        upload_form = Booking(request.POST, request.FILES)

        if upload_form.is_valid():
            upload_form.save()
            return HttpResponse("Ticket Added Successfully")
        else:
            print(upload_form.errors)    
            return HttpResponse(upload_form.errors)

    else:
        upload_form = Booking()

    return render(request, "two/upload.html", {"upload_form": upload_form})

def ticket_list(request):
    tickets=Tickets.objects.all()
    return render(request, 'two/ticket_list.html',{'tickets': tickets})

def update_ticket(request, ticket_id):
    ticket_id = int(ticket_id)
    try:
        ticket_sel = Tickets.objects.get(id=ticket_id)
    except Tickets.DoesNotExist:
        return redirect('ticket_list')
    upload_form = Booking(request.POST or None, instance=ticket_sel)
    if upload_form.is_valid():
        upload_form.save()
        return redirect('ticket_list')
    return render(request, 'two/upload.html',{'upload_form': upload_form})

def delete_ticket(request, ticket_id):
    ticket_id = int(ticket_id)
    try:
        ticket_sel = Tickets.objects.get(id=ticket_id)
    except Tickets.DoesNotExist:
        return redirect('ticket_list')
    ticket_sel.delete()
    return redirect('ticket_list')

def customer_list(request):
    customers = Customer.objects.all()
    return render(request, "two/customer_list.html", {"customers": customers})

@csrf_exempt
def customer_upload(request):
    if request.method == "POST":
        form = CustomerForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect("customer_list")
    else:
        form = CustomerForm()

    return render(request, "two/customer_upload.html", {"form": form})

@csrf_exempt
def customer_update(request, id):
    customer = Customer.objects.get(CustomerId=id)

    if request.method == "POST":
        form = CustomerForm(request.POST, instance=customer)
        if form.is_valid():
            form.save()
            return redirect("customer_list")
    else:
        form = CustomerForm(instance=customer)

    return render(request, "two/customer_upload.html", {"form": form})


def customer_delete(request, id):
    try:
        customer = Customer.objects.get(CustomerId=id)
    except Customer.DoesNotExist:
        return redirect("customer_list")

    customer.delete()
    return redirect("customer_list")


def agent_list(request):
    agents = Agent.objects.all()
    return render(request, "two/agent_list.html", {"agents": agents})

@csrf_exempt
def agent_upload(request):
    if request.method == "POST":
        form = AgentForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('agent_list')
    else:
        form = AgentForm()

    return render(request, 'two/agent_upload.html', {'form': form})

@csrf_exempt
def update_agent(request, agent_id):
    try:
        agent = Agent.objects.get(id=agent_id)
    except Agent.DoesNotExist:
        return redirect('agent_list')

    form = AgentForm(request.POST or None, instance=agent)

    if form.is_valid():
        form.save()
        return redirect('agent_list')

    return render(request, 'two/agent_upload.html', {'form': form})

@csrf_exempt
def delete_agent(request, agent_id):
    try:
        agent = Agent.objects.get(id=agent_id)
    except Agent.DoesNotExist:
        return redirect('agent_list')

    agent.delete()
    return redirect('agent_list') 





def api_offer(request):
    if request.method=='GET':
        username = request.GET.get("username")
        password = request.GET.get("password")
        user = authenticate(username=username, password=password)
        if user:
            return JsonResponse({'message': ' login done.'})
            return JsonResponse({'success': True}, status=200)
        else:
           return JsonResponse({'success': False,'message': 'Invalid username or password'}, status=400)
    else:
        return JsonResponse({'message': ' only GET method is allowed.'},status= 405)


def dashboard(request):
    return render(request,'dashboard.html')   

def tickets(request):
    return render(request,'tickets.html')   

from django.shortcuts import render
from .models import Tickets, Customer, Agent

def dashboard(request):
    all_tickets = Tickets.objects.all()

    context = {
        'tickets': all_tickets.select_related('customer_id', 'assigned').order_by('-id')[:10],
        'total_tickets_count': all_tickets.count(),
        'progress_tickets_count': all_tickets.filter(status__iexact='In Progress').count(),
        'todo_tickets_count': all_tickets.filter(status__iexact='To Do').count(),
        'completed_tickets_count': all_tickets.filter(status__iexact='Completed').count(),
        'rejected_tickets_count': all_tickets.filter(status__iexact='Rejected').count(),
    }

    return render(request, 'dashboard.html', context)

def tickets(request):
    all_tickets = Tickets.objects.all()
    return render(request, 'tickets.html', {'tickets': all_tickets})  

def report(request):
    return render(request, "report.html")



def generate_report_view(request):
    context = {
        "customers": Customer.objects.all(),
        "agents": Agent.objects.all(),
        "action_type": None,
        "selected_status": None,
        "data_items": None,
        "start_date": "",
        "end_date": "",
    }

    if request.method == "POST":
        status = request.POST.get("status")
        action = request.POST.get("action")
        customer = request.POST.get("customer")
        priority = request.POST.get("priority")
        agent = request.POST.get("agent")
        start_date = request.POST.get("start_date")
        end_date = request.POST.get("end_date")

        context["selected_status"] = status
        context["start_date"] = start_date
        context["end_date"] = end_date
        context["selected_priority"] = priority
        context["selected_customer"] = customer
        context["selected_agent"] = agent

        # Start with all tickets
        queryset = Tickets.objects.all()

        # Apply filters only if they are selected
        if status:
            queryset = queryset.filter(status=status)

        if customer:
            queryset = queryset.filter(customer_id=customer)

        if agent:
            queryset = queryset.filter(assigned_id=agent)

        if priority:
            queryset = queryset.filter(priority=priority)

        if start_date:
            queryset = queryset.filter(created_At__date__gte=start_date)

        if end_date:
            queryset = queryset.filter(created_At__date__lte=end_date)

        # Show report
        if action == "show":
            context["action_type"] = "show"
            context["data_items"] = queryset
            return render(request, "report.html", context)

        # Export report
        elif action == "export":
            output = BytesIO()
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Report"

            headers = [
                "ID",
                "Priority",
                "Status",
                "Date",
            ]

            for col, header in enumerate(headers, 1):
                cell = worksheet.cell(row=1, column=col)
                cell.value = header
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            row = 2

            for ticket in queryset:
                worksheet.cell(row=row, column=1).value = ticket.id
                worksheet.cell(row=row, column=2).value = ticket.priority
                worksheet.cell(row=row, column=3).value = ticket.status
                worksheet.cell(row=row, column=4).value = ticket.created_At.date()
                row += 1

            workbook.save(output)
            output.seek(0)

            response = HttpResponse(
                output.read(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            filename = f"{status}_Report.xlsx" if status else "Report.xlsx"

            response["Content-Disposition"] = (
                f'attachment; filename="{filename}"'
            )

            return response

    return render(request, "report.html", context)